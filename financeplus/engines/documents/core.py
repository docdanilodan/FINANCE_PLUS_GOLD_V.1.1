from __future__ import annotations

from dataclasses import dataclass, asdict
from io import BytesIO
from pathlib import Path
from typing import Any


@dataclass
class DocumentAnalysis:
    filename: str
    category: str = "Da verificare"
    confidence: float = 0.0
    company_name: str | None = None
    document_year: int | None = None
    suggested_name: str | None = None
    extracted_text: str = ""
    status: str = "IMPLEMENTATO"
    warnings: list[str] | None = None

    def to_dict(self) -> dict[str, Any]:
        out = asdict(self)
        out["warnings"] = list(self.warnings or [])
        return out


def extract_text_bytes(filename: str, data: bytes) -> tuple[str, list[str]]:
    """Extract text locally where possible; never invent OCR output."""
    ext = Path(filename).suffix.lower()
    warnings: list[str] = []
    if not data:
        return "", ["File vuoto."]
    try:
        if ext in {".txt", ".csv", ".md", ".json", ".xml", ".ofx", ".qif"}:
            return data.decode("utf-8", errors="replace"), warnings
        if ext == ".pdf":
            from pypdf import PdfReader
            text = "\n".join((p.extract_text() or "") for p in PdfReader(BytesIO(data)).pages)
            if not text.strip(): warnings.append("PDF senza testo estraibile: RICHIEDE OCR/IDP o verifica umana.")
            return text, warnings
        if ext == ".docx":
            from docx import Document
            doc = Document(BytesIO(data)); return "\n".join(p.text for p in doc.paragraphs), warnings
        if ext in {".xlsx", ".xls"}:
            import pandas as pd
            xl = pd.ExcelFile(BytesIO(data)); chunks = []
            for sheet in xl.sheet_names:
                frame = pd.read_excel(BytesIO(data), sheet_name=sheet, header=None)
                chunks.append(frame.to_csv(index=False, header=False))
            return "\n".join(chunks), warnings
        if ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(BytesIO(data)); chunks: list[str] = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text: chunks.append(shape.text)
            return "\n".join(chunks), warnings
    except Exception as exc:
        warnings.append(f"Estrazione locale non riuscita: {exc}")
        return "", warnings
    warnings.append(f"Formato {ext or '(senza estensione)'} archiviabile ma non estratto localmente.")
    return "", warnings


def classify_document(filename: str, data: bytes, pasted_text: str = "") -> DocumentAnalysis:
    """Use the verified repository Document AI when available; fail safe otherwise."""
    text, warnings = extract_text_bytes(filename, data)
    source_text = (text + "\n" + (pasted_text or "")).strip()
    try:
        from document_ai import classify_text, suggested_name
        classification = classify_text((filename + "\n" + source_text)[:120000])
        result = DocumentAnalysis(
            filename=filename,
            category=str(getattr(classification, "category", "Da verificare")),
            confidence=float(getattr(classification, "confidence", 0.0) or 0.0),
            company_name=getattr(classification, "company_name", None),
            document_year=getattr(classification, "document_year", None),
            extracted_text=source_text,
            warnings=warnings,
        )
        try:
            result.suggested_name = suggested_name(classification, filename)
        except Exception as exc:
            result.warnings = list(result.warnings or []) + [f"Naming automatico non disponibile: {exc}"]
        if not source_text: result.status = "RICHIEDE OCR/IDP"
        elif result.confidence < 0.70: result.status = "DA TESTARE / REVISIONE UMANA"
        return result
    except Exception as exc:
        return DocumentAnalysis(filename=filename, extracted_text=source_text, status="RICHIEDE SORGENTE ORIGINALE", warnings=warnings + [f"Classifier repository non disponibile: {exc}"])
